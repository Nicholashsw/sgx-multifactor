"""
PIT adapter — swaps the current-constituent panel for the survivorship-corrected one,
then re-runs the Part 0-6 attribution framework unchanged.

STATUS: written against the SGX3_2.xlsm schema recovered from the SGX3 chat. NOT YET
EXECUTED against the real workbook — no numbers from this file are verified until it
has run on the actual data.

Expected schema (SGX3_2.xlsm):
  PriceData_PIT    : Date, Ticker, PX_LAST, TOT_RET_IDX, MKT_CAP_USDmn, PX_TO_BOOK, DIV_YIELD, PX_VOLUME
  Fundamentals_PIT : Date, Ticker, PE_RATIO, ROE, DEBT_TO_EQ, PX_TO_BOOK_Q
  Universe_PIT     : Ticker, First, Last          (membership window, quarter-ends)
  Classification2  : Ticker, Sector               (or Classification)
  Benchmark        : Date, STI

Falls back to SGX2_v3.xlsm layout (PricePanel / Membership / Classification2) if the
_PIT sheets are absent.
"""
import numpy as np
import pandas as pd
from openpyxl import load_workbook

# exit-type conventions for the delisting return, run as BOUNDS not as a point estimate
DELIST_CONVENTIONS = {
    "truncate":  None,     # drop the final month  -> upper bound (optimistic)
    "minus100":  -1.00,    # assign -100%          -> lower bound (pessimistic)
    "classified": "table", # per-name from exit_types.csv if supplied
}


def read_sheet(wb, name, cols):
    ws = wb[name]
    n = len(cols)
    data = [r[:n] for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]
    df = pd.DataFrame(data, columns=cols)
    if "Date" in cols:
        df["Date"] = pd.to_datetime(df["Date"], errors="coerce")
        df = df.dropna(subset=["Date"])
    for c in cols:
        if c in ("Date", "Ticker", "Sector"):
            continue
        df[c] = pd.to_numeric(
            df[c].apply(lambda v: np.nan if isinstance(v, str) and "#" in v else v),
            errors="coerce")
    return df


def truncate_padding(px, price_col="PX_LAST", tol=1e-9):
    """Bloomberg pads a dead ticker's last price forward to the request end-date.
    Cut each ticker at its last genuine price change so dead names leave the panel."""
    keep, log = [], []
    for t, g in px.sort_values("Date").groupby("Ticker"):
        v = g[price_col].to_numpy(float)
        moved = np.where(np.abs(np.diff(v)) > tol)[0]
        if len(moved) == 0:
            last = len(g) - 1
        else:
            last = moved[-1] + 1
        dropped = len(g) - 1 - last
        if dropped > 0:
            log.append({"Ticker": t, "padded_months": dropped,
                        "true_last": g["Date"].iloc[last].date()})
        keep.append(g.iloc[:last + 1])
    return pd.concat(keep, ignore_index=True), pd.DataFrame(log)


def build_pit_panel(wb_path, embargo=4, delist="truncate", exit_types=None,
                    start="2008-01-31"):
    """Returns a panel with the SAME column names the attribution engine expects."""
    wb = load_workbook(wb_path, data_only=True, read_only=True)
    sheets = set(wb.sheetnames)

    if "PriceData_PIT" in sheets:
        px = read_sheet(wb, "PriceData_PIT",
                        ["Date", "Ticker", "PX_LAST", "TOT_RETURN", "MKT_CAP_USDmn",
                         "PB_mo", "DY_mo", "PX_VOLUME"])
        fn = read_sheet(wb, "Fundamentals_PIT",
                        ["Date", "Ticker", "PE_q", "ROE_q", "DEBT_TO_EQ", "PB_q"])
        mem = pd.DataFrame([{"Ticker": r[0], "First": pd.to_datetime(r[1]),
                             "Last": pd.to_datetime(r[2])}
                            for r in wb["Universe_PIT"].iter_rows(min_row=2, values_only=True)
                            if r[0] and "SP" in str(r[0])])
    elif "PricePanel" in sheets:      # SGX2_v3 layout
        px = read_sheet(wb, "PricePanel",
                        ["Date", "Ticker", "PX_LAST", "TOT_RETURN", "MKT_CAP_USDmn",
                         "PB_mo", "DY_mo", "PX_VOLUME", "BVPS"])
        fn = read_sheet(wb, "Fundamentals",
                        ["Date", "Ticker", "PE_q", "ROE_q", "PB_q"])
        msheet = "MembershipQ" if "MembershipQ" in sheets else "Membership"
        raw = read_sheet(wb, msheet, ["Date", "Ticker"])
        mem = (raw.groupby("Ticker")["Date"].agg(First="min", Last="max").reset_index())
    else:
        raise KeyError(f"no recognised price sheet in {wb_path}: {sorted(sheets)}")

    csheet = "Classification2" if "Classification2" in sheets else "Classification"
    cls = read_sheet(wb, csheet, ["Ticker", "Sector"])[["Ticker", "Sector"]]

    bm = read_sheet(wb, "Benchmark", ["Date", "STI"])
    bm["Date"] = bm["Date"] + pd.offsets.MonthEnd(0)
    bm["STI_Ret"] = bm["STI"].sort_index().pct_change()

    for d in (px, fn):
        d["Date"] = d["Date"] + pd.offsets.MonthEnd(0)
    px = px.sort_values(["Ticker", "Date"]).drop_duplicates(["Ticker", "Date"])
    fn = fn.sort_values(["Ticker", "Date"]).drop_duplicates(["Ticker", "Date"])

    # 1. cut Bloomberg's post-delisting padding
    px, padlog = truncate_padding(px)
    print(f"padding truncated on {len(padlog)} tickers "
          f"({padlog.padded_months.sum() if len(padlog) else 0} name-months removed)")

    # 2. reporting embargo on the fiscal-dated fundamentals
    fn["AvailFrom"] = fn["Date"] + pd.offsets.MonthEnd(embargo)
    parts = []
    for t, p in px.groupby("Ticker"):
        f = fn[fn.Ticker == t][["AvailFrom", "PE_q", "ROE_q", "PB_q"]].sort_values("AvailFrom")
        p = p.sort_values("Date")
        parts.append(p.assign(PE_q=np.nan, ROE_q=np.nan, PB_q=np.nan) if f.empty else
                     pd.merge_asof(p, f, left_on="Date", right_on="AvailFrom",
                                   direction="backward",
                                   tolerance=pd.Timedelta("200D")).drop(columns=["AvailFrom"]))
    p = pd.concat(parts, ignore_index=True).merge(cls, on="Ticker", how="left")

    # 3. returns from the total-return index
    p = p.sort_values(["Ticker", "Date"]).reset_index(drop=True)
    p["Ret"] = p.groupby("Ticker")["TOT_RETURN"].pct_change()

    # 4. delisting return on the terminal month of every name that dies before sample end
    sample_end = p["Date"].max()
    last_obs = p.groupby("Ticker")["Date"].transform("max")
    dead = (last_obs < sample_end - pd.offsets.MonthEnd(1))
    terminal = dead & (p["Date"] == last_obs)
    n_dead = p.loc[terminal, "Ticker"].nunique()
    if delist == "minus100":
        p.loc[terminal, "Ret"] = -1.00
    elif delist == "truncate":
        p = p[~terminal]
    elif delist == "classified":
        if exit_types is None:
            raise ValueError("delist='classified' needs exit_types (Ticker, ExitType)")
        et = exit_types.set_index("Ticker")["ExitType"].to_dict()
        for i in p.index[terminal]:
            kind = et.get(p.at[i, "Ticker"], "unknown")
            if kind in ("distress", "insolvent", "liquidation"):
                p.at[i, "Ret"] = -1.00
            elif kind in ("acquisition", "merger", "take-private"):
                pass                       # last quote is the payout, keep as measured
            else:
                p.at[i, "Ret"] = np.nan    # unclassified: exclude, do not guess
    print(f"delisting convention '{delist}' applied to {n_dead} terminal name-months")

    # 5. point-in-time membership filter
    mem["First"] = pd.to_datetime(mem["First"]) + pd.offsets.MonthEnd(-1)
    mem["Last"] = pd.to_datetime(mem["Last"]) + pd.offsets.MonthEnd(0)
    p = p.merge(mem, on="Ticker", how="left")
    before = len(p)
    p = p[(p["Date"] >= p["First"]) & (p["Date"] <= p["Last"])]
    print(f"PIT membership filter: {before:,} -> {len(p):,} name-months")

    # 6. derived factors, identical definitions to the current-constituent run
    p = p.sort_values(["Ticker", "Date"]).reset_index(drop=True)
    p["MOM_12_1"] = p.groupby("Ticker")["TOT_RETURN"].pct_change(11).shift(1)
    p["MOM_6_1"] = p.groupby("Ticker")["TOT_RETURN"].pct_change(5).shift(1)
    p["BM"] = 1.0 / p["PB_mo"].clip(lower=0.05)
    p["DY"] = p["DY_mo"]
    p["EY"] = 1.0 / p["PE_q"].where(p["PE_q"] > 0)
    p["ROE"] = p["ROE_q"]
    p["SIZE"] = -np.log(p["MKT_CAP_USDmn"].clip(lower=1))
    p["ADV_USDmn"] = p["PX_VOLUME"] * p["PX_LAST"] / 21 / 1e6 / 1.35
    p["LIQ"] = np.log(p["ADV_USDmn"].clip(lower=0.01))
    p = p.merge(bm[["Date", "STI_Ret"]], on="Date", how="left")

    p = p[p["Date"] >= start]
    print(f"final PIT panel: {len(p):,} rows | {p.Ticker.nunique()} names | "
          f"{p.Date.min().date()}-{p.Date.max().date()} | "
          f"median {p.groupby('Date').size().median():.0f} names/month")
    return p, padlog


if __name__ == "__main__":
    import sys
    from attribution import WEIGHTS, run, stats_block, score, RF_M, START  # engine reuse

    wbp = sys.argv[1] if len(sys.argv) > 1 else "SGX3_2.xlsm"
    for conv in ("truncate", "minus100"):
        print("\n" + "=" * 90)
        print(f"PIT RUN — delisting convention: {conv}")
        print("=" * 90)
        pit, _ = build_pit_panel(wbp, embargo=4, delist=conv)
        res, contrib = run(pit, WEIGHTS, cost_bps=0.0)
        res = res.set_index("Date")
        print(pd.DataFrame([
            stats_block(res["Long"], res["Bench"], f"Long-only top-6 [{conv}]"),
            stats_block(res["LS"],   res["Bench"], f"Long-short 6v6 [{conv}]"),
            stats_block(res["Bench"], None,        "Equal-weight PIT universe"),
        ]).round(3).to_string(index=False))

        # THE test the current-constituent panel could not run:
        # does survivorship correction act on the long leg or the short leg?
        c = contrib.merge(pit[["Date", "Ticker", "Sector"]], on=["Date", "Ticker"], how="left")
        yrs = c.Date.nunique() / 12
        legs = c.groupby("Leg")["c"].sum() / yrs * 100
        print(f"\nleg-level P&L (annualised pp): " +
              "  ".join(f"{k}={v:+.2f}" for k, v in legs.items()))
        res.to_csv(f"pit_returns_{conv}.csv")
        contrib.to_csv(f"pit_contrib_{conv}.csv", index=False)
