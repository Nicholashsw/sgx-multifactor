"""
SGX multi-factor. pipeline.py

Builds the monthly macro panel and the 30-name Bloomberg stock panel.
Run: python pipeline.py

Inputs (same folder):
  bloombergapi4_fixed.xlsm   Bloomberg workbook (PriceData, Fundamentals, Classification)
  sti sgx.xlsx, sti sgx 2.xlsx, bfeus factor (1).xlsx   CEIC macro exports
  SPX INDEX.xlsx, VIX Index.xlsx, DXY Curncy.xlsx, FDTR Index.xlsx,
  SPY US Equity.xlsx         Bloomberg daily market series

Outputs:
  macro_monthly.csv, panel_stock.csv

Design rules, non-negotiable:
  1. Real data only. There is no synthetic fallback anywhere in this file.
     If an input is missing or a required column is absent, the run stops
     with an explicit error. Earlier versions of this pipeline contained a
     silent fallback to a seeded random generator; that path is removed.
  2. Returns and momentum are computed from TOT_RETURN (total-return index),
     not PX_LAST. Price-only returns discard the dividend, which is a third
     of total return in this market and biases the dividend-yield factor.
  3. Fundamentals are merged as-of with backward direction only, so no
     future fiscal data reaches an earlier month. A further 4-month
     reporting embargo is applied downstream in factor_model.py.
"""

import datetime
import sys
from functools import reduce
from pathlib import Path

import numpy as np
import pandas as pd

DATA_DIR = Path(".")

REQUIRED_PANEL_COLS = ["PX_LAST", "TOT_RETURN", "MKT_CAP", "PX_TO_BOOK",
                       "DIV_YIELD", "PX_VOLUME"]


def fail(msg: str) -> None:
    """Stop the run loudly. Silent degradation is how the last pipeline died."""
    sys.exit(f"PIPELINE ERROR: {msg}")


def is_date_cell(v) -> bool:
    return isinstance(v, (datetime.datetime, datetime.date, pd.Timestamp))


def date_block_start(first_col: pd.Series, name: str) -> int:
    """Position of the first date row in an export whose header rows sit above
    the data. Every row from there to the end must be a date; a non-date row
    inside the block means the export is malformed, so the run stops."""
    flags = first_col.map(is_date_cell).to_numpy()
    if not flags.any():
        fail(f"{name}: no date rows found")
    start = int(flags.argmax())
    bad = np.flatnonzero(~flags[start:])
    if len(bad):
        row = start + int(bad[0])
        fail(f"{name}: {len(bad)} non-date row(s) inside the data block, "
             f"first at row {row}: {first_col.iloc[row]!r}")
    return start


# Bloomberg daily market files

def parse_bbg_daily(path: Path, col: str) -> pd.DataFrame:
    """Bloomberg daily export: a few header rows (Security, Start Date, ...)
    above a Date | Value block."""
    df = pd.read_excel(path, header=None)
    start = date_block_start(df.iloc[:, 0], path.name)
    data = df.iloc[start:, :2]
    out = pd.DataFrame({"Date": pd.to_datetime(data.iloc[:, 0]),
                        col: pd.to_numeric(data.iloc[:, 1], errors="coerce")})
    return out.set_index("Date").sort_index()


def build_market_factors() -> pd.DataFrame:
    series = {
        "SPX INDEX.xlsx": "SPX",
        "VIX Index.xlsx": "VIX",
        "DXY Curncy.xlsx": "DXY",
        "FDTR Index.xlsx": "FDTR",
        "SPY US Equity.xlsx": "SPY",
    }
    frames = []
    for fname, col in series.items():
        path = DATA_DIR / fname
        if not path.exists():
            fail(f"required market file missing: {fname}")
        df = parse_bbg_daily(path, col)
        frames.append(df.resample("ME").last())
        print(f"  Loaded {fname}: {len(df)} daily rows")

    mkt = frames[0]
    for frame in frames[1:]:
        mkt = mkt.join(frame, how="outer")

    mkt["SPX_Ret_1M"] = mkt["SPX"].pct_change()
    mkt["SPX_Ret_3M"] = mkt["SPX"].pct_change(3)
    mkt["SPX_Ret_12M"] = mkt["SPX"].pct_change(12)
    mkt["VIX_MA3"] = mkt["VIX"].rolling(3).mean()
    mkt["DXY_Ret_3M"] = mkt["DXY"].pct_change(3)
    mkt["Regime_RiskOn"] = ((mkt["VIX"] < 20) & (mkt["SPX_Ret_3M"] > 0)).astype(int)
    mkt["Regime_VIX_High"] = (mkt["VIX"] > 25).astype(int)
    mkt.index.name = "Date"
    return mkt.reset_index()


# CEIC macro files

def parse_ceic(path: Path):
    if not path.exists():
        fail(f"required CEIC file missing: {path.name}")
    df = pd.read_excel(path, sheet_name=0, header=None)
    # CEIC exports put a series-metadata block (Region, Frequency, Unit, ...)
    # above the dated observations.
    ds = date_block_start(df.iloc[:, 0], path.name)
    data = df.iloc[ds:].copy()
    data.columns = ["Date"] + list(range(1, len(data.columns)))
    data["Date"] = pd.to_datetime(data["Date"])
    data = data.set_index("Date").sort_index()
    for c in data.columns:
        data[c] = pd.to_numeric(data[c], errors="coerce")
    return data


def build_macro() -> pd.DataFrame:
    sg = parse_ceic(DATA_DIR / "sti sgx.xlsx").resample("ME").last()
    sg.columns = [
        "SG_IPI", "SG_PMI", "SG_CPI_YoY", "SG_NODX_Elec",
        "SG_Visitors_CN", "SG_Visitors_Total", "USDSGD",
        "SG_GDP_YoY", "SG_SORA_3M", "SG_NEER",
        "SG_SGS_2Y", "SG_SGS_10Y", "SG_M2_YoY", "SG_CoreCPI_YoY",
    ][:len(sg.columns)]

    sg2 = parse_ceic(DATA_DIR / "sti sgx 2.xlsx").resample("ME").last()
    sg2.columns = ["SG_GDP_GovtConsump", "SG_GDP_PrivConsump",
                   "SG_GDP_GFCF", "CN_IPI_YoY"][:len(sg2.columns)]

    gl = parse_ceic(DATA_DIR / "bfeus factor (1).xlsx").resample("ME").last()
    gl.columns = [
        "CN_PPI_YoY", "CN_PMI", "JP_Tankan",
        "SG_CoreCPI2", "SG_CPI2", "SG_GDP2",
        "US_M2_YoY", "US_Unemployment", "US_ISM_PMI",
        "US_CoreCPI_YoY", "US_CPI_YoY",
    ][:len(gl.columns)]

    macro = pd.concat([sg, sg2, gl], axis=1, sort=True)
    macro["SG_YieldSpread"] = macro["SG_SGS_10Y"] - macro["SG_SGS_2Y"]
    macro["SG_PMI_Change"] = macro["SG_PMI"].diff()
    macro["CN_PMI_Change"] = macro["CN_PMI"].diff()
    macro["SG_SORA_Change"] = macro["SG_SORA_3M"].diff()
    macro["US_ISM_Change"] = macro["US_ISM_PMI"].diff()
    macro["SG_CPI_Surprise"] = (macro["SG_CPI_YoY"]
                                - macro["SG_CPI_YoY"].rolling(3).mean().shift(1))
    macro["Regime_SGRates_Rising"] = (macro["SG_SORA_Change"] > 0).astype(int)
    macro["Regime_SGGrowth"] = (macro["SG_PMI"] > 50).astype(int)
    macro["Regime_ChinaGrowth"] = (macro["CN_PMI"] > 50).astype(int)
    macro.index.name = "Date"
    return macro.reset_index()


# Bloomberg workbook

def load_bloomberg_stock_panel() -> pd.DataFrame:
    """
    Reads all metric blocks from bloombergapi4_fixed.xlsm.
    PriceData: six monthly metrics in horizontal layout, one 3-column block
    per ticker (Date | Value | blank). Fundamentals: three quarterly metrics,
    merged as-of into the monthly panel.
    """
    from openpyxl import load_workbook

    workbook_path = DATA_DIR / "bloombergapi4_fixed.xlsm"
    if not workbook_path.exists():
        fail("bloombergapi4_fixed.xlsm not found. There is no fallback; "
             "export the workbook before running.")

    print("  Loading bloombergapi4_fixed.xlsm with openpyxl...")
    wb = load_workbook(str(workbook_path), keep_vba=True, data_only=True)

    def read_block(ws, block_start_col, metric_name,
                   stride=3, data_start=4, max_rows=310, max_tickers=30):
        frames = []
        for t_idx in range(max_tickers):
            date_col = block_start_col + t_idx * stride
            val_col = date_col + 1
            ticker = ws.cell(row=2, column=date_col).value
            if not ticker or "SP Equity" not in str(ticker):
                break
            ticker = str(ticker).strip()
            dates, vals = [], []
            for r in range(data_start, data_start + max_rows):
                d = ws.cell(row=r, column=date_col).value
                v = ws.cell(row=r, column=val_col).value
                if d is None:
                    break
                where = f"{ws.title}!R{r}C{date_col} ({ticker}, {metric_name})"
                if is_date_cell(d):
                    d = pd.Timestamp(d)
                else:
                    try:
                        d = pd.to_datetime(d)
                    except (ValueError, TypeError):
                        fail(f"unreadable date {d!r} at {where}")
                if v is None or (isinstance(v, str) and "#" in v):
                    v = np.nan      # empty cell or Bloomberg error code (#N/A ...)
                else:
                    try:
                        v = float(v)
                    except (ValueError, TypeError):
                        fail(f"non-numeric value {v!r} at {where}")
                dates.append(d)
                vals.append(v)
            if dates:
                frames.append(pd.DataFrame(
                    {"Date": dates, metric_name: vals, "Ticker": ticker}))
        return pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()

    ws_p = wb["PriceData"]
    price_blocks = [("PX_LAST", 1), ("TOT_RETURN", 93), ("MKT_CAP", 185),
                    ("PX_TO_BOOK", 277), ("DIV_YIELD", 369), ("PX_VOLUME", 461)]
    price_dfs = []
    for metric, start_col in price_blocks:
        df = read_block(ws_p, start_col, metric, max_rows=310)
        if df.empty:
            fail(f"PriceData block '{metric}' parsed to zero rows")
        print(f"    {metric}: {df['Ticker'].nunique()} tickers, "
              f"{df[metric].notna().sum()} valid values")
        price_dfs.append(df)

    panel = reduce(lambda l, r: pd.merge(l, r, on=["Date", "Ticker"],
                                         how="outer"), price_dfs)

    ws_f = wb["Fundamentals"]
    fund_blocks = [("PE_RATIO", 1), ("ROE", 93), ("PX_TO_BOOK_Q", 185)]
    fund_dfs = []
    for metric, start_col in fund_blocks:
        df = read_block(ws_f, start_col, metric, max_rows=110)
        if not df.empty:
            print(f"    {metric}: {df['Ticker'].nunique()} tickers, "
                  f"{df[metric].notna().sum()} valid values")
            fund_dfs.append(df)

    if fund_dfs:
        fund_panel = reduce(lambda l, r: pd.merge(l, r, on=["Date", "Ticker"],
                                                  how="outer"), fund_dfs)
        fund_panel = fund_panel.sort_values(["Ticker", "Date"])
        merged_parts = []
        for ticker in panel["Ticker"].unique():
            p = panel[panel["Ticker"] == ticker].sort_values("Date")
            f = fund_panel[fund_panel["Ticker"] == ticker].sort_values("Date")
            if f.empty:
                merged_parts.append(p)
            else:
                merged_parts.append(pd.merge_asof(
                    p, f.drop(columns=["Ticker"]), on="Date",
                    direction="backward"))
        panel = pd.concat(merged_parts, ignore_index=True)

    ws_c = wb["Classification"]
    class_rows = []
    for row in ws_c.iter_rows(min_row=2, values_only=True):
        if row[0] and "SP Equity" in str(row[0]):
            class_rows.append({"Ticker": row[0], "Sector": row[2],
                               "Industry": row[3]})
    if class_rows:
        panel = pd.merge(panel, pd.DataFrame(class_rows), on="Ticker",
                         how="left")

    missing = [c for c in REQUIRED_PANEL_COLS if c not in panel.columns]
    if missing:
        fail(f"panel missing required columns after load: {missing}. "
             "Fix the workbook export; do not paper over this.")

    print(f"  Real panel: {len(panel):,} rows, {panel['Ticker'].nunique()} tickers")
    return panel.sort_values(["Ticker", "Date"]).reset_index(drop=True)


def main():
    print("=== Building macro panel ===")
    mkt = build_market_factors()
    macro_ceic = build_macro()
    macro = pd.merge(mkt, macro_ceic, on="Date", how="outer").sort_values("Date")
    macro["Date"] = macro["Date"] + pd.offsets.MonthEnd(0)
    macro.to_csv(DATA_DIR / "macro_monthly.csv", index=False)
    print(f"  macro_monthly.csv: {len(macro)} rows x {len(macro.columns)} cols")

    print("\n=== Building stock panel ===")
    stock = load_bloomberg_stock_panel()
    stock = stock.rename(columns={"MKT_CAP": "MKT_CAP_SGDmn"})

    stock["Date"] = pd.to_datetime(stock["Date"]) + pd.offsets.MonthEnd(0)
    stock = stock.sort_values(["Ticker", "Date"]).reset_index(drop=True)

    # Returns and momentum from the total-return index. Never PX_LAST.
    g = stock.groupby("Ticker")
    stock["Return_1M"] = g["TOT_RETURN"].pct_change(1)
    stock["Return_3M"] = g["TOT_RETURN"].pct_change(3)
    stock["MOM_12_1"] = g["TOT_RETURN"].pct_change(12).groupby(stock["Ticker"]).shift(1)
    stock["MOM_6_1"] = g["TOT_RETURN"].pct_change(6).groupby(stock["Ticker"]).shift(1)
    stock["MOM_3_1"] = g["TOT_RETURN"].pct_change(3).groupby(stock["Ticker"]).shift(1)
    stock["LogMktCap"] = np.log(stock["MKT_CAP_SGDmn"].clip(lower=1))
    stock["BM_Ratio"] = 1.0 / stock["PX_TO_BOOK"].clip(lower=0.1)
    stock["LogVolume"] = np.log(stock["PX_VOLUME"].clip(lower=1))
    stock["Amihud"] = stock["Return_1M"].abs() / stock["PX_VOLUME"].clip(lower=1e6)

    stock = pd.merge(stock, macro, on="Date", how="left")
    stock.to_csv(DATA_DIR / "panel_stock.csv", index=False)
    print(f"  panel_stock.csv: {len(stock):,} rows | "
          f"{stock['Ticker'].nunique()} tickers")
    print(f"  Date: {stock['Date'].min().date()} to {stock['Date'].max().date()}")
    print("\nDone. Next: python factor_model.py")


if __name__ == "__main__":
    main()
