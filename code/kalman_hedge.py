"""
Dynamic hedge ratios: Kalman filter vs static OLS, rolling OLS and EWMA beta.

Test 1  beta hedge   short beta * market against each name and the equal-weight book
Test 2  pairs        Kalman spread trading within GICS industry clusters

The tests take a date x name total-return index frame and a market return
series (or one market series per name, for a leave-one-out proxy), so the
same code runs on the monthly Bloomberg panel and the daily Compustat panel.

Train up to TRAIN_END (Kalman noise parameters estimated here, then frozen),
holdout from TEST_START onwards, touched once.

Usage:  python kalman_hedge.py bloombergapi4_fixed.xlsm [out_dir]
"""
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import openpyxl
import statsmodels.api as sm
from scipy.optimize import minimize
from statsmodels.tsa.stattools import coint

try:
    from numba import njit
except ImportError:          # plain Python is fine for monthly, slow for daily
    def njit(f):
        return f

TRAIN_END = "2019-11-30"
TEST_START = "2019-12-01"
COST_PAIR = 0.0030          # one-way per leg, SGX cash
COST_HEDGE = 0.0005         # one-way on hedge notional changes, STI futures
DELTA_FIXED = 1e-4          # forced state noise for the robustness variant

FREQ = {
    "monthly": dict(roll=36, hl=12, nw=6, ann=12, min_train=60, min_test=24, burn=12),
    "daily": dict(roll=252, hl=63, nw=10, ann=252, min_train=756, min_test=250, burn=60),
}

CLUSTERS = {
    "Banks": ["DBS", "OCBC", "UOB"],
    "Industrial REITs": ["CLAR", "FLT", "MINT", "MLT"],
    "Retail/Div REITs": ["CICT", "FCT", "MPACT"],
    "RE developers": ["CIT", "HKL", "UOL"],
}


# ---------------------------------------------------------------- monthly loader

def load_bloomberg_monthly(path):
    """Total return index (gross dividends) for the 30 STI names, plus the STI price index."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb["PriceData"].iter_rows(values_only=True))
    tickers, block = rows[1], 92          # TR index block starts at column 92
    tr = {}
    for j in range(block, block + 90, 3):
        name = tickers[j].replace(" SP Equity", "")
        s = {r[j]: r[j + 1] for r in rows[3:]
             if r[j] is not None and isinstance(r[j + 1], (int, float))}
        tr[name] = pd.Series(s, dtype=float)
    tr = pd.DataFrame(tr)
    tr.index = pd.to_datetime(tr.index).to_period("M").to_timestamp("M")

    bench = list(wb["Benchmark"].iter_rows(values_only=True))[1:]
    sti = pd.Series({r[0]: r[1] for r in bench if r[0] is not None and r[1]}, dtype=float)
    sti.index = pd.to_datetime(sti.index).to_period("M").to_timestamp("M")
    return tr.sort_index(), sti.sort_index()


# ---------------------------------------------------------------- kalman

@njit
def _filter(y, x, q0, q1, r, a0, a1, p0):
    """Two-state random-walk regression y_t = a_t + b_t x_t + e_t.
    Returns the state known at t-1 (used for trading at t), the innovation and its variance."""
    n = len(y)
    prior = np.full((n, 2), np.nan)
    innov = np.full(n, np.nan)
    svar = np.full(n, np.nan)
    a, b = a0, a1
    p00, p01, p11 = p0, 0.0, p0
    for t in range(n):
        p00 += q0
        p11 += q1
        prior[t, 0] = a
        prior[t, 1] = b
        if np.isnan(y[t]) or np.isnan(x[t]):
            continue
        xt = x[t]
        e = y[t] - a - b * xt
        h0 = p00 + p01 * xt                # P @ [1, x]
        h1 = p01 + p11 * xt
        s = h0 + h1 * xt + r
        k0, k1 = h0 / s, h1 / s
        a += k0 * e
        b += k1 * e
        p00 -= k0 * h0
        p01 -= k0 * h1
        p11 -= k1 * h1
        innov[t] = e
        svar[t] = s
    return prior, innov, svar


def kalman_filter(y, x, q, r, a0, p0=1.0):
    return _filter(np.asarray(y, float), np.asarray(x, float),
                   float(q[0]), float(q[1]), float(r), float(a0[0]), float(a0[1]), p0)


def fit_kalman(y, x, burn):
    """Prediction-error MLE for the two state variances and the observation variance."""
    ols = sm.OLS(y, sm.add_constant(x), missing="drop").fit()
    a0, r0 = np.asarray(ols.params), ols.resid.var()

    def nll(theta):
        q, r = np.exp(theta[:2]), np.exp(theta[2])
        _, e, s = kalman_filter(y, x, q, r, a0)
        e, s = e[burn:], s[burn:]
        ok = ~np.isnan(e)
        return 0.5 * np.sum(np.log(s[ok]) + e[ok] ** 2 / s[ok])

    x0 = np.r_[np.log(np.full(2, r0 * 1e-3)), np.log(r0)]
    res = minimize(nll, x0, method="Nelder-Mead", options={"maxiter": 4000, "xatol": 1e-4})
    return np.exp(res.x[:2]), np.exp(res.x[2]), a0


# ---------------------------------------------------------------- helpers

def rolling_beta(y, x, window):
    return (y.rolling(window).cov(x) / x.rolling(window).var()).shift(1)


def ewma_beta(y, x, halflife):
    ew = lambda s: s.ewm(halflife=halflife).mean()
    cov = ew(y * x) - ew(y) * ew(x)
    return (cov / x.ewm(halflife=halflife).var(bias=True)).shift(1)


def nw_t(x, lags):
    x = pd.Series(x).dropna()
    if len(x) < 12:
        return np.nan
    m = sm.OLS(x.values, np.ones(len(x))).fit(cov_type="HAC", cov_kwds={"maxlags": lags})
    return float(m.tvalues[0])


def dm_test(e1, e2, lags):
    """Diebold-Mariano on squared hedge errors. Negative t means e1 is smaller."""
    return nw_t(pd.Series(e1) ** 2 - pd.Series(e2) ** 2, lags)


def sharpe(x, ann):
    x = pd.Series(x).dropna()
    return x.mean() / x.std() * np.sqrt(ann) if x.std() > 0 else np.nan


def hold_every(b, step):
    """Keep the hedge ratio fixed between rebalances (step=5 on daily data is weekly)."""
    if step <= 1:
        return b
    keep = np.arange(len(b)) % step == 0
    return b.where(keep).ffill()


# ---------------------------------------------------------------- test 1

def beta_hedge_test(ret, market, freq="monthly", rebalance=1):
    """market: one return series, or a DataFrame with a column per name
    (plus 'EW book') for a leave-one-out market proxy."""
    cfg = FREQ[freq]
    targets = dict(ret)
    targets["EW book"] = ret.mean(axis=1, skipna=True)
    rows, paths = [], {}

    for name, y in targets.items():
        m = market[name] if isinstance(market, pd.DataFrame) else market
        df = pd.concat([y, m], axis=1, keys=["y", "m"]).dropna()
        tr = df[df.index <= TRAIN_END]
        oos = df.index >= TEST_START
        if len(tr) < cfg["min_train"] or oos.sum() < cfg["min_test"]:
            continue

        b_static = sm.OLS(tr.y, sm.add_constant(tr.m)).fit().params["m"]
        q, r, a0 = fit_kalman(tr.y.values, tr.m.values, cfg["burn"])
        prior, _, _ = kalman_filter(df.y.values, df.m.values, q, r, a0)

        betas = pd.DataFrame({
            "Static OLS": b_static,
            "Rolling": rolling_beta(df.y, df.m, cfg["roll"]),
            "EWMA": ewma_beta(df.y, df.m, cfg["hl"]),
            "Kalman": prior[:, 1],
        }, index=df.index).apply(hold_every, step=rebalance)

        unhedged = df.y[oos]
        row = {"name": name, "n_oos": int(oos.sum()), "q_beta": q[1], "r": r}
        errs = {}
        for col in betas:
            b = betas[col]
            cost = b.diff().abs().fillna(0) * COST_HEDGE
            errs[col] = (df.y - b * df.m - cost)[oos]
            row[col] = 1 - errs[col].var() / unhedged.var()
        lags = cfg["nw"]
        row["DM Kalman vs Rolling"] = dm_test(errs["Kalman"], errs["Rolling"], lags)
        row["DM Kalman vs EWMA"] = dm_test(errs["Kalman"], errs["EWMA"], lags)
        row["DM Kalman vs Static"] = dm_test(errs["Kalman"], errs["Static OLS"], lags)
        row["beta_sd_kalman"] = betas["Kalman"][oos].std()
        row["beta_sd_rolling"] = betas["Rolling"][oos].std()
        rows.append(row)
        paths[name] = betas

    return pd.DataFrame(rows).set_index("name"), paths


# ---------------------------------------------------------------- test 2

def trade_spread(z, entry):
    """+1 long spread below -entry, -1 short above +entry, flat when z crosses zero.
    The position decided at t earns the return over t+1."""
    pos, cur = np.zeros(len(z)), 0
    for t, v in enumerate(z.values):
        if not np.isnan(v):
            if cur == 0:
                cur = 1 if v < -entry else (-1 if v > entry else 0)
            elif (cur == 1 and v >= 0) or (cur == -1 and v <= 0):
                cur = 0
        pos[t] = cur
    return pd.Series(pos, index=z.index)


def pair_returns(pos, b, ry, rx):
    """Long 1 of y, short b of x, scaled to unit gross; costs on traded notional."""
    gross = 1 + b.abs()
    wy, wx = pos / gross, -pos * b / gross
    gret = wy.shift(1) * ry + wx.shift(1) * rx
    traded = wy.diff().abs().fillna(wy.abs()) + wx.diff().abs().fillna(wx.abs())
    return gret.fillna(0) - traded * COST_PAIR


def pairs_test(tri, freq="monthly", entries=(1.0, 1.5, 2.0)):
    cfg = FREQ[freq]
    logp, ret = np.log(tri), tri.pct_change(fill_method=None)
    results, series, diag = [], {}, []

    for cluster, names in CLUSTERS.items():
        names = [n for n in names if n in tri.columns]
        for i in range(len(names)):
            for j in range(i + 1, len(names)):
                yn, xn = names[i], names[j]
                df = pd.concat([logp[yn], logp[xn], ret[yn], ret[xn]], axis=1,
                               keys=["y", "x", "ry", "rx"]).dropna()
                tr = df[df.index <= TRAIN_END]
                if len(tr) < cfg["min_train"] or (df.index >= TEST_START).sum() < cfg["min_test"]:
                    continue
                pair = f"{yn}/{xn}"

                # static Engle-Granger, fitted on train only
                eg = sm.OLS(tr.y, sm.add_constant(tr.x)).fit()
                a_s, b_s = eg.params
                z_static = (df.y - a_s - b_s * df.x - eg.resid.mean()) / eg.resid.std()
                rho = sm.OLS(eg.resid.values[1:], eg.resid.values[:-1]).fit().params[0]

                # rolling OLS hedge ratio
                w = cfg["roll"]
                b_r = df.y.rolling(w).cov(df.x) / df.x.rolling(w).var()
                a_r = df.y.rolling(w).mean() - b_r * df.x.rolling(w).mean()
                res_r = df.y - a_r.shift(1) - b_r.shift(1) * df.x
                z_roll = (res_r - res_r.rolling(w).mean().shift(1)) / res_r.rolling(w).std().shift(1)

                # kalman with MLE noise, and with forced dynamics
                q, r, a0 = fit_kalman(tr.y.values, tr.x.values, cfg["burn"])
                prior, e, s = kalman_filter(df.y.values, df.x.values, q, r, a0)
                d = DELTA_FIXED / (1 - DELTA_FIXED)
                pf, ef, sf = kalman_filter(df.y.values, df.x.values, [d, d], eg.resid.var(), a0)
                oos = df.index >= TEST_START

                specs = {
                    "Static EG": (z_static, pd.Series(b_s, index=df.index)),
                    "Rolling": (z_roll, b_r.shift(1)),
                    "Kalman MLE": (pd.Series(e / np.sqrt(s), index=df.index),
                                   pd.Series(prior[:, 1], index=df.index)),
                    "Kalman fixed delta": (pd.Series(ef / np.sqrt(sf), index=df.index),
                                           pd.Series(pf[:, 1], index=df.index)),
                }
                diag.append({
                    "pair": pair, "cluster": cluster,
                    "eg_pvalue": coint(tr.y, tr.x)[1],
                    "half_life": -np.log(2) / np.log(rho) if 0 < rho < 1 else np.inf,
                    "beta_static": b_s,
                    "kalman_mle_beta_oos_sd": specs["Kalman MLE"][1][oos].std(),
                    "kalman_fixed_beta_oos_sd": specs["Kalman fixed delta"][1][oos].std(),
                })

                for method, (z, b) in specs.items():
                    for k in entries:
                        pos = trade_spread(z, k)
                        pr = pair_returns(pos, b.fillna(b_s), df.ry, df.rx)
                        series[(pair, method, k)] = pr
                        for period, mask in (("train", df.index <= TRAIN_END), ("oos", oos)):
                            x = pr[mask]
                            results.append({
                                "pair": pair, "cluster": cluster, "method": method,
                                "entry": k, "period": period, "mean": x.mean(),
                                "t_nw": nw_t(x, cfg["nw"]), "sharpe": sharpe(x, cfg["ann"]),
                                "trades": int((pos[mask].diff().abs() > 0).sum()),
                            })
    return pd.DataFrame(results), series, pd.DataFrame(diag)


def pairs_portfolio(series, freq="monthly"):
    cfg, rows = FREQ[freq], []
    for m, k in sorted({(m, k) for _, m, k in series}):
        port = pd.concat([v for (_, mm, kk), v in series.items() if (mm, kk) == (m, k)], axis=1).mean(axis=1)
        for period, mask in (("train", port.index <= TRAIN_END), ("oos", port.index >= TEST_START)):
            x = port[mask]
            rows.append({"method": m, "entry": k, "period": period, "mean": x.mean(),
                         "t_nw": nw_t(x, cfg["nw"]), "sharpe": sharpe(x, cfg["ann"])})
    return pd.DataFrame(rows)


# ---------------------------------------------------------------- main

def run(tri, market, freq, out, tag, rebalance=1):
    out.mkdir(parents=True, exist_ok=True)
    hedge, _ = beta_hedge_test(tri.pct_change(fill_method=None), market, freq, rebalance)
    pairs, series, diag = pairs_test(tri, freq)
    port = pairs_portfolio(series, freq)
    hedge.to_csv(out / f"hedge_{tag}.csv")
    pairs.to_csv(out / f"pairs_{tag}.csv", index=False)
    port.to_csv(out / f"pairs_portfolio_{tag}.csv", index=False)
    diag.to_csv(out / f"pair_diagnostics_{tag}.csv", index=False)
    return hedge, pairs, port, diag


if __name__ == "__main__":
    tri, sti = load_bloomberg_monthly(sys.argv[1])
    out = Path(sys.argv[2] if len(sys.argv) > 2 else "results/kalman")
    mkt = sti.pct_change().reindex(tri.index)
    hedge, pairs, port, diag = run(tri, mkt, "monthly", out, "monthly_30_sti")

    names = hedge.drop("EW book")
    print(names[["Static OLS", "Rolling", "EWMA", "Kalman"]].mean().round(4))
    print(port.pivot_table(index=["method", "entry"], columns="period", values="t_nw").round(2))
